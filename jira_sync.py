#!/usr/bin/env python3
"""Replica worklogs do Jira do cliente para o Jira interno.

Versão 3:
- preserva os comentários fixos exigidos pelo cliente;
- gera descrições internas no formato solicitado pelo RH;
- usa o status real da issue do cliente sem presumir conclusão;
- trabalha em modo de simulação por padrão (só grava com --apply);
- evita duplicidade por worklog de origem, sem pular um dia inteiro;
- bloqueia lançamentos que fariam o total diário ultrapassar 8 horas.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

import requests
from requests.auth import HTTPBasicAuth

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover - dependência já usada pela versão original
    def load_dotenv(*_args: Any, **_kwargs: Any) -> bool:
        return False


load_dotenv(Path(__file__).with_name(".env"))

SRC_BASE_URL = os.getenv("SRC_BASE_URL", "https://ti-segurosunimed-soa.atlassian.net").rstrip("/")
SRC_EMAIL = os.getenv("SRC_EMAIL")
SRC_API_TOKEN = os.getenv("SRC_API_TOKEN")
SRC_ACCOUNT_ID = os.getenv("SRC_ACCOUNT_ID", "")

DST_BASE_URL = os.getenv("DST_BASE_URL", "https://jira-uxorit.atlassian.net").rstrip("/")
DST_EMAIL = os.getenv("DST_EMAIL")
DST_API_TOKEN = os.getenv("DST_API_TOKEN")
DST_ACCOUNT_ID = os.getenv("DST_ACCOUNT_ID", "")

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-5-nano")
OPENAI_RESPONSES_URL = "https://api.openai.com/v1/responses"
AI_BATCH_SIZE = 30

WEEKS_BACK = int(os.getenv("WEEKS_BACK", "2"))
DAILY_LIMIT_SECONDS = 8 * 60 * 60
SYNC_PROPERTY_KEY = "jira_sync_v2_source"

# A ordem importa: categorias específicas vêm antes das genéricas.
TASK_MAPPING = [
    (["code review", "revisão de código", "revisao de codigo"], "UN-33"),
    (["daily", "dailly", "dayli"], "UN-12"),
    (["desenvolvimento", "development", "implementação", "implementacao"], "UN-14"),
    (["reunião", "reuniao", "meeting", "reunio"], "UN-24"),
    (["sustentação", "sustentacao", "sustenta", "hypercare"], "UN-22"),
    (["documentação", "documentacao", "document", "definição", "definicao", "arquitetura"], "UN-26"),
    (["apoio", "apoio time", "support"], "UN-27"),
    (["teste", "testes", "homologação", "homologacao"], "UN-31"),
]

DEDICATED_SUMMARY_MAPPING = {
    "daily": "UN-12",
    "reuniao": "UN-24",
    "apoio ao time": "UN-27",
    "code review": "UN-33",
    "code review desenvolvimento": "UN-33",
    "code review homologacao": "UN-33",
    "teste": "UN-31",
    "testes": "UN-31",
}

GENERIC_PREFIXES = {
    "produtor", "consumidor", "desenvolvimento", "incidente", "debito tecnico",
    "gestao", "sustentacao", "documentacao", "teste", "testes", "reuniao",
}

GENERIC_TITLES = {
    "daily", "reuniao", "desenvolvimento", "documentacao", "sustentacao",
    "apoio", "apoio ao time", "teste", "testes", "hypercare", "log rotate",
    "implantacao de robo", "ajustes em homologacao", "code review",
    "code review desenvolvimento", "code review homologacao",
    "incidente", "debito tecnico", "gestao",
}

SHORT_ACTIVITY_TITLES = {
    "hypercare", "log rotate", "implantacao de robo", "ajustes em homologacao",
    "code review", "code review desenvolvimento", "code review homologacao",
}

AI_INSTRUCTIONS = """Você revisa apontamentos de horas para um RH brasileiro.
Para cada item, devolva duas frases curtas, objetivas e profissionais em português do Brasil:
`what_done`, descrevendo a atividade realizada, e `deliverable`, descrevendo o avanço ou resultado.
Use somente fatos presentes no item. O status é controlado pelo script: não o altere nem o repita.
Quando o status não for Finalizado, jamais afirme que algo foi concluído, corrigido, entregue ou está
pronto. Preserve siglas úteis, mas remova chaves do Jira, datas, horas, nomes de pessoas e expressões
como 'registro de'. Evite textos genéricos. Retorne exatamente um resultado para cada id recebido e
siga estritamente o esquema JSON."""


def normalize(text: Any) -> str:
    value = unicodedata.normalize("NFKD", str(text or ""))
    return " ".join(
        "".join(ch for ch in value if not unicodedata.combining(ch)).lower().split()
    )


def is_fixed_daily(item: dict[str, Any]) -> bool:
    return normalize(item.get("key")) == "tirpa-560" and normalize(item.get("comment")) == "reuniao"


def parse_jira_datetime(value: str | None) -> datetime | None:
    if not value:
        return None
    value = value.strip()
    if value.endswith("Z"):
        value = value[:-1] + "+00:00"
    elif len(value) >= 5 and value[-5] in ("+", "-") and value[-3] != ":":
        value = value[:-2] + ":" + value[-2:]
    return datetime.fromisoformat(value)


def seconds_to_hm(seconds: int) -> str:
    hours, remainder = divmod(int(seconds), 3600)
    minutes = remainder // 60
    return f"{hours}h{minutes:02d}m"


def extract_comment(raw: Any) -> str:
    if not raw:
        return ""
    if isinstance(raw, str):
        return " ".join(raw.split())
    if isinstance(raw, list):
        return " ".join(filter(None, (extract_comment(item) for item in raw))).strip()
    if isinstance(raw, dict):
        own = str(raw.get("text", "")).strip()
        children = extract_comment(raw.get("content", []))
        return " ".join(part for part in (own, children) if part).strip()
    return str(raw).strip()


def clean_issue_summary(text: Any) -> str:
    """Limpa prefixos técnicos sem apagar nomes úteis como [AC19 ...]."""
    value = " ".join(str(text or "").split()).strip(" -–—")
    match = re.match(r"^\[([^]]+)]\s*(.*)$", value)
    if not match:
        return value
    prefix, remainder = match.group(1).strip(), match.group(2).strip(" -–—")
    if normalize(prefix) in GENERIC_PREFIXES:
        return remainder or prefix
    return f"{prefix} — {remainder}" if remainder else prefix


def simplify_context_label(text: Any, compact: bool = True) -> str:
    """Reduz ruído estrutural de títulos sem produzir conteúdo novo."""
    raw = " ".join(str(text or "").split()).strip(" -–—")
    bracket = re.match(r"^\[([^]]+)]\s*(.*)$", raw)
    if bracket:
        scope, remainder = bracket.group(1).strip(), bracket.group(2).strip(" -–—")
        if normalize(scope) not in GENERIC_PREFIXES:
            incident = re.search(r"\b(falha|erro)\b", normalize(remainder))
            code = re.search(r"\b[A-Z]{2,}\d+\b", scope)
            if incident and code:
                kind = "falha" if incident.group(1) == "falha" else "erro"
                return f"{kind} no {code.group(0)}"
            if compact:
                scope = re.sub(r"\s*\([^)]*\)\s*", " ", scope).strip()
                return re.sub(r"\s*[-–—]\s*", " ", scope).strip()

    value = clean_issue_summary(raw)
    process_match = re.match(r"(?i)^processo\s+de\s+(.+)$", value)
    if compact and process_match and re.search(r"\s[-–—]\s", process_match.group(1)):
        value = process_match.group(1)
    value = re.sub(r"(?i)\brobô\s+RPA\b", "", value)
    value = re.sub(r"\s*[-–—]\s*", " ", value)
    value = " ".join(value.split()).strip()
    return value


def match_category(text: Any) -> str | None:
    value = normalize(text)
    for keywords, issue_key in TASK_MAPPING:
        if any(normalize(keyword) in value for keyword in keywords):
            return issue_key
    return None


def map_to_target_issue(comment: Any, summary: Any) -> str | None:
    dedicated = DEDICATED_SUMMARY_MAPPING.get(normalize(clean_issue_summary(summary)))
    return dedicated or match_category(comment) or match_category(summary)


def _is_generic(text: Any) -> bool:
    value = normalize(text).strip(" -–—()")
    return value in GENERIC_TITLES


def _activity_with_context(issue: str, context: str) -> str:
    if normalize(issue) not in SHORT_ACTIVITY_TITLES:
        return ""
    return f"{issue} — {context}" if context else issue


def _join_contexts(contexts: Iterable[str], limit: int = 2) -> str:
    unique: list[str] = []
    for raw in contexts:
        value = simplify_context_label(raw, compact=True)
        if not value or _is_generic(value) or normalize(value) in {normalize(x) for x in unique}:
            continue
        unique.append(value)
        if len(unique) == limit:
            break
    if not unique:
        return ""
    if len(unique) == 1:
        return unique[0]
    return f"{unique[0]} e {unique[1]}"


def build_internal_description(
    comment: Any,
    summary: Any,
    parent_summary: Any = "",
    robot_name: Any = "",
    day_contexts: Iterable[str] = (),
) -> str:
    """Cria uma descrição objetiva usando somente contexto existente no Jira."""
    category = normalize(comment)
    issue = simplify_context_label(summary, compact=False)
    parent = simplify_context_label(parent_summary, compact=True)
    robot = simplify_context_label(robot_name, compact=True)
    same_day = _join_contexts(day_contexts)
    context = robot or (parent if parent and not _is_generic(parent) else "")
    specific_issue = issue if issue and not _is_generic(issue) else ""
    short_activity = _activity_with_context(issue, context or same_day)

    if "daily" in category or normalize(issue) == "daily":
        subject = same_day or context
        return (
            f"Reunião diária — alinhamento sobre {subject}"
            if subject else "Reunião diária de acompanhamento das atividades do cliente"
        )

    if any(term in category for term in ("reuniao", "meeting")):
        subject = specific_issue or same_day or context
        return (
            f"Reunião de alinhamento — {subject}"
            if subject else "Reunião de alinhamento das atividades do cliente"
        )

    if "code review" in category or "revisao de codigo" in category:
        subject = short_activity or specific_issue or context or same_day
        return f"Revisão de código — {subject}" if subject else "Revisão de código da automação"

    if any(term in category for term in ("teste", "homologacao")):
        if normalize(issue) == "hypercare":
            subject = context or same_day
            return f"Validação em hypercare — {subject}" if subject else "Validação em período de hypercare"
        subject = short_activity or specific_issue or context or same_day
        return f"Testes e validação — {subject}" if subject else "Testes e validação da automação"

    if any(term in category for term in ("sustentacao", "hypercare", "support")):
        if normalize(issue) == "hypercare":
            subject = context or same_day
            return (
                f"Sustentação em hypercare do {subject}"
                if subject else "Sustentação em período de hypercare"
            )
        subject = short_activity or specific_issue or context or same_day
        return f"Sustentação — {subject}" if subject else "Sustentação das automações do cliente"

    if any(term in category for term in ("documentacao", "definicao", "arquitetura")):
        subject = short_activity or specific_issue or context or same_day
        return f"Documentação — {subject}" if subject else "Elaboração e atualização de documentação técnica"

    if any(term in category for term in ("desenvolvimento", "development", "implementacao")):
        if short_activity:
            return f"Desenvolvimento — {short_activity}"
        if specific_issue:
            if re.match(r"^(criar|analisar|ajustar|implementar|configurar|preencher|modelar)\b", normalize(specific_issue)):
                return f"Desenvolvimento — {specific_issue}"
            return f"Desenvolvimento do {specific_issue}"
        subject = context or same_day
        return f"Desenvolvimento — {subject}" if subject else "Desenvolvimento e evolução da automação"

    if any(term in category for term in ("apoio", "support")):
        subject = short_activity or specific_issue or same_day or context
        return f"Apoio ao time — {subject}" if subject else "Apoio ao time em demandas do cliente"

    return specific_issue or context or same_day or str(comment or "Atividade do cliente").strip()


def resolve_ai_choice(
    cli_value: bool | None,
    input_fn: Any = input,
    is_tty: bool | None = None,
) -> bool:
    """Resolve a opção de IA sem bloquear execuções automatizadas."""
    if cli_value is not None:
        return bool(cli_value)
    if is_tty is None:
        is_tty = sys.stdin.isatty()
    if not is_tty:
        return False
    answer = str(input_fn("Usar IA para melhorar as descrições? [s/N] ")).strip().lower()
    return answer in {"s", "sim", "y", "yes"}


def _response_output_text(payload: dict[str, Any]) -> str:
    direct = payload.get("output_text")
    if isinstance(direct, str) and direct.strip():
        return direct
    for output in payload.get("output", []):
        if not isinstance(output, dict):
            continue
        for content in output.get("content", []):
            if not isinstance(content, dict):
                continue
            if content.get("type") == "output_text" and isinstance(content.get("text"), str):
                return content["text"]
    raise ValueError("A resposta da IA não contém texto estruturado.")


def _valid_ai_description(value: Any) -> bool:
    text = " ".join(str(value or "").split())
    if not 12 <= len(text) <= 220:
        return False
    if re.search(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b", text):
        return False
    if re.search(r"\b[A-Z][A-Z0-9]{1,}-\d+\b", text):
        return False
    if "registro de" in normalize(text):
        return False
    return True


def _claims_completion(value: Any) -> bool:
    text = normalize(value)
    return bool(re.search(
        r"\b(concluid[oa]s?|finalizad[oa]s?|corrigid[oa]s?|entregue|entregues|pronto|pronta|"
        r"resolvid[oa]s?)\b",
        text,
    ))


def _ai_item_id(item: dict[str, Any]) -> str:
    return str(item.get("id") or str(item.get("source_marker", "")).rsplit("|", 1)[-1])


def _ai_request_body(items: list[dict[str, Any]], model: str) -> dict[str, Any]:
    safe_items = [{
        "id": _ai_item_id(item),
        "categoria": item.get("comment", ""),
        "titulo_origem": _record_summary(item),
        "titulo_pai": _record_parent(item),
        "atividade_sem_ia": item.get("what_done", ""),
        "resultado_sem_ia": item.get("deliverable", ""),
        "status_controlado_pelo_script": item.get("status_text", "Em andamento"),
    } for item in items]
    return {
        "model": model,
        "instructions": AI_INSTRUCTIONS,
        "input": json.dumps({"items": safe_items}, ensure_ascii=False),
        "text": {
            "format": {
                "type": "json_schema",
                "name": "jira_timesheet_descriptions",
                "strict": True,
                "schema": {
                    "type": "object",
                    "properties": {
                        "entries": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "id": {"type": "string"},
                                    "what_done": {"type": "string"},
                                    "deliverable": {"type": "string"},
                                },
                                "required": ["id", "what_done", "deliverable"],
                                "additionalProperties": False,
                            },
                        }
                    },
                    "required": ["entries"],
                    "additionalProperties": False,
                },
            }
        },
        "reasoning": {"effort": "minimal"},
        "store": False,
        "max_output_tokens": max(1600, len(items) * 120),
    }


def enhance_plan_with_ai(
    plan: Iterable[dict[str, Any]],
    api_key: str,
    model: str = OPENAI_MODEL,
    session: Any = None,
) -> list[dict[str, Any]]:
    """Melhora descrições em lote e mantém o resultado local em caso de falha."""
    result = []
    for original in plan:
        item = dict(original)
        item["description_without_ai"] = item["description"]
        if is_fixed_daily(item):
            item["description"] = "Daily"
            item["description_without_ai"] = "Daily"
            item["ai_status"] = "fixed_daily"
        result.append(item)

    candidates = [
        item for item in result
        if item.get("status") == "ready" and not is_fixed_daily(item)
    ]
    if not api_key:
        for item in candidates:
            item["ai_status"] = "fallback_missing_key"
        return result

    client = session or requests.Session()
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    for index in range(0, len(candidates), AI_BATCH_SIZE):
        batch = candidates[index:index + AI_BATCH_SIZE]
        try:
            response = client.post(
                OPENAI_RESPONSES_URL,
                headers=headers,
                json=_ai_request_body(batch, model),
                timeout=90,
            )
            response.raise_for_status()
            parsed = json.loads(_response_output_text(response.json()))
            rows = parsed.get("entries", []) if isinstance(parsed, dict) else []
            entries = {
                str(row.get("id")): {
                    "what_done": " ".join(str(row.get("what_done", "")).split()),
                    "deliverable": " ".join(str(row.get("deliverable", "")).split()),
                }
                for row in rows if isinstance(row, dict) and row.get("id")
            }
            for item in batch:
                proposed = entries.get(_ai_item_id(item), {})
                what_done = proposed.get("what_done", "")
                deliverable = proposed.get("deliverable", "")
                completion_is_allowed = item.get("status_text") == "Finalizado"
                valid = (
                    _valid_ai_description(what_done)
                    and _valid_ai_description(deliverable)
                    and (
                        completion_is_allowed
                        or not (_claims_completion(what_done) or _claims_completion(deliverable))
                    )
                )
                if valid:
                    item["what_done"] = what_done
                    item["deliverable"] = deliverable
                    item["description"] = format_rh_description(
                        what_done,
                        item.get("status_text", "Em andamento"),
                        deliverable,
                    )
                    item["ai_status"] = "enhanced"
                else:
                    item["ai_status"] = "fallback_invalid"
        except (requests.RequestException, RuntimeError, ValueError, TypeError, json.JSONDecodeError):
            for item in batch:
                item["ai_status"] = "fallback_error"
    return result


def source_marker(source_base_url: str, source_worklog_id: Any) -> str:
    return f"{source_base_url.rstrip('/')}|{source_worklog_id}"


def _record_date(item: dict[str, Any]) -> str:
    return str(item.get("started", ""))[:10]


def _record_summary(item: dict[str, Any]) -> str:
    return str(item.get("summary") or item.get("issue_summary") or "")


def _record_parent(item: dict[str, Any]) -> str:
    return str(item.get("parentSummary") or item.get("parent_summary") or "")


def derive_status(item: dict[str, Any]) -> str:
    """Traduz o estado real da issue de origem sem presumir conclusão."""
    source_status = " ".join(str(item.get("sourceStatus") or "").split())
    category = normalize(item.get("sourceStatusCategory") or "")
    resolution = " ".join(str(item.get("sourceResolution") or "").split())
    normalized_status = normalize(source_status)

    if (
        category in {"done", "concluido", "finalizado"}
        or normalized_status in {"done", "concluido", "finalizado", "fechado", "closed"}
        or resolution
    ):
        return "Finalizado"

    if source_status:
        if normalized_status in {
            "a fazer", "aberto", "backlog", "novo", "open", "selected for development",
            "to do",
        }:
            return "Em andamento"
        return source_status
    return "Em andamento"


def _result_context(item: dict[str, Any]) -> str:
    summary = clean_issue_summary(_record_summary(item))
    parent = clean_issue_summary(_record_parent(item))

    documentation = re.fullmatch(r"(?i)analisar documenta[cç][aã]o\s*\((.+)\)", summary)
    if documentation:
        summary = documentation.group(1).strip()

    context = summary
    if not context or _is_generic(context):
        context = parent
    context = simplify_context_label(context, compact=True)
    return context or "atividade registrada no Jira do cliente"


def build_deliverable_result(item: dict[str, Any], what_done: str) -> str:
    """Gera um resultado conservador, sem declarar encerramento da atividade."""
    category = normalize(extract_comment(item.get("comment")))
    context = _result_context(item)
    if "—" in what_done:
        activity_context = what_done.rsplit("—", 1)[1].strip()
        if activity_context and len(activity_context) < len(context):
            context = activity_context
    templates = {
        "apoio": "Apoio prestado na atividade “{context}”.",
        "code review": "Revisão técnica registrada na atividade “{context}”.",
        "desenvolvimento": "Evolução da atividade “{context}”.",
        "documentacao": "Avanço na documentação da atividade “{context}”.",
        "reuniao": "Alinhamento realizado na atividade “{context}”.",
        "sustentacao": "Avanço na tratativa da atividade “{context}”.",
        "teste": "Validações registradas na atividade “{context}”.",
    }
    template = templates.get(category, "Avanço registrado na atividade “{context}”.")
    return template.format(context=context)


def format_rh_description(what_done: str, status_text: str, deliverable: str) -> str:
    return (
        f"O que foi feito: {what_done}\n"
        f"Status: {status_text}\n"
        f"Entregável / Resultado: {deliverable}"
    )


def build_sync_plan(
    source_records: Iterable[dict[str, Any]],
    existing_source_markers: set[str] | None = None,
    source_base_url: str = SRC_BASE_URL,
) -> list[dict[str, Any]]:
    records = [dict(item) for item in source_records]
    existing = existing_source_markers or set()
    contexts_by_day: dict[str, list[str]] = defaultdict(list)

    for item in records:
        raw_summary = _record_summary(item)
        raw_parent = _record_parent(item)
        summary = clean_issue_summary(raw_summary)
        parent = clean_issue_summary(raw_parent)
        parent_is_useful = bool(parent) and not _is_generic(parent)
        if _is_generic(summary):
            candidate = raw_parent
        elif len(summary) > 64 and parent_is_useful:
            candidate = raw_parent
        else:
            candidate = raw_summary
        if candidate and not _is_generic(clean_issue_summary(candidate)):
            contexts_by_day[_record_date(item)].append(candidate)

    plan: list[dict[str, Any]] = []
    for item in sorted(records, key=lambda x: (str(x.get("started", "")), str(x.get("id", "")))):
        comment = extract_comment(item.get("comment"))
        summary = _record_summary(item)
        parent = _record_parent(item)
        target = map_to_target_issue(comment, summary)
        marker = source_marker(source_base_url, item.get("id", ""))
        own = {
            normalize(simplify_context_label(summary, compact=True)),
            normalize(simplify_context_label(parent, compact=True)),
        }
        day_contexts = [
            value for value in contexts_by_day.get(_record_date(item), [])
            if normalize(simplify_context_label(value, compact=True)) not in own
        ]
        what_done = build_internal_description(
            comment=comment,
            summary=summary,
            parent_summary=parent,
            robot_name=item.get("robotName", ""),
            day_contexts=day_contexts,
        )
        status_text = derive_status(item)
        deliverable = build_deliverable_result({**item, "comment": comment}, what_done)
        description = format_rh_description(what_done, status_text, deliverable)
        if is_fixed_daily({**item, "comment": comment}):
            what_done = "Daily"
            status_text = ""
            deliverable = ""
            description = "Daily"
        if marker in existing:
            status = "already_synced"
        elif not target:
            status = "unmapped"
        else:
            status = "ready"
        plan.append({
            **item,
            "comment": comment,
            "target_issue": target,
            "description": description,
            "what_done": what_done,
            "status_text": status_text,
            "deliverable": deliverable,
            "source_marker": marker,
            "status": status,
        })
    return plan


def get_period_range(weeks_back: int = WEEKS_BACK) -> tuple[date, date]:
    today = datetime.now().date()
    monday = today - timedelta(days=today.weekday())
    return monday - timedelta(weeks=max(weeks_back, 1) - 1), monday + timedelta(days=6)


def apply_daily_limit(
    plan: Iterable[dict[str, Any]],
    existing_daily_seconds: dict[str, int] | None = None,
    limit_seconds: int = DAILY_LIMIT_SECONDS,
) -> list[dict[str, Any]]:
    """Calcula os totais diários e bloqueia o que ultrapassaria o limite."""
    existing = {
        str(day): int(seconds)
        for day, seconds in (existing_daily_seconds or {}).items()
    }
    running = defaultdict(int, existing)
    result: list[dict[str, Any]] = []

    for original in plan:
        item = dict(original)
        day = _record_date(item)
        current = running[day]
        projected = current
        if item.get("status") == "ready":
            projected = current + int(item.get("seconds") or item.get("timeSpentSeconds") or 0)
            if projected > limit_seconds:
                item["status"] = "probable_duplicate"
            else:
                running[day] = projected
                current = projected
        item["daily_existing_seconds"] = existing.get(day, 0)
        item["daily_projected_seconds"] = projected
        item["daily_total_seconds"] = current
        result.append(item)
    return result


def _session(email: str | None, token: str | None) -> requests.Session:
    session = requests.Session()
    if email and token:
        session.auth = HTTPBasicAuth(email, token)
    session.headers.update({"Accept": "application/json"})
    return session


def _account_id(session: requests.Session, base_url: str) -> str:
    response = session.get(f"{base_url}/rest/api/3/myself", timeout=30)
    response.raise_for_status()
    return response.json()["accountId"]


def fetch_source_records(start: date, end: date) -> list[dict[str, Any]]:
    if not (SRC_EMAIL and SRC_API_TOKEN):
        raise RuntimeError("Defina SRC_EMAIL e SRC_API_TOKEN para consultar o Jira de origem.")
    session = _session(SRC_EMAIL, SRC_API_TOKEN)
    account_id = SRC_ACCOUNT_ID or _account_id(session, SRC_BASE_URL)
    since = int(datetime(start.year, start.month, start.day, tzinfo=timezone.utc).timestamp() * 1000)
    ids: set[str] = set()
    params = {"since": since}
    while True:
        response = session.get(
            f"{SRC_BASE_URL}/rest/api/3/worklog/updated", params=params, timeout=60
        )
        response.raise_for_status()
        payload = response.json()
        ids.update(str(item["worklogId"]) for item in payload.get("values", []) if item.get("worklogId"))
        if payload.get("lastPage", True):
            break
        params["since"] = payload.get("until", params["since"])

    worklogs: list[dict[str, Any]] = []
    ids_list = sorted(ids)
    for index in range(0, len(ids_list), 1000):
        response = session.post(
            f"{SRC_BASE_URL}/rest/api/3/worklog/list",
            json={"ids": ids_list[index:index + 1000]}, timeout=60,
        )
        response.raise_for_status()
        worklogs.extend(response.json())

    issue_cache: dict[str, dict[str, Any]] = {}
    records: dict[str, dict[str, Any]] = {}
    for worklog in worklogs:
        worklog_id = str(worklog.get("id", ""))
        started = parse_jira_datetime(worklog.get("started"))
        if not worklog_id or not started or not (start <= started.date() <= end):
            continue
        if worklog.get("author", {}).get("accountId") != account_id:
            continue
        issue_id = str(worklog.get("issueId", ""))
        if issue_id not in issue_cache:
            response = session.get(
                f"{SRC_BASE_URL}/rest/api/3/issue/{issue_id}",
                params={"fields": "summary,parent,status,resolution"}, timeout=30,
            )
            response.raise_for_status()
            issue_cache[issue_id] = response.json()
        issue = issue_cache[issue_id]
        fields = issue.get("fields", {})
        parent = fields.get("parent") or {}
        source_status = fields.get("status") or {}
        source_status_category = source_status.get("statusCategory") or {}
        source_resolution = fields.get("resolution") or {}
        candidate = {
            "id": worklog_id,
            "updated": worklog.get("updated", ""),
            "key": issue.get("key", issue_id),
            "summary": fields.get("summary", ""),
            "parentKey": parent.get("key", ""),
            "parentSummary": (parent.get("fields") or {}).get("summary", ""),
            "sourceStatus": source_status.get("name", ""),
            "sourceStatusCategory": source_status_category.get("key", ""),
            "sourceResolution": source_resolution.get("name", ""),
            "started": worklog.get("started", ""),
            "seconds": int(worklog.get("timeSpentSeconds", 0)),
            "comment": extract_comment(worklog.get("comment")),
        }
        previous = records.get(worklog_id)
        if not previous or str(candidate["updated"]) > str(previous.get("updated", "")):
            records[worklog_id] = candidate
    return list(records.values())


def get_destination_state(
    session: requests.Session,
    account_id: str,
    start: date,
    end: date,
) -> tuple[set[str], dict[str, int]]:
    """Obtém marcadores e horas do usuário em todas as issues visíveis do destino."""
    markers: set[str] = set()
    daily_seconds: dict[str, int] = defaultdict(int)
    since = int(datetime(start.year, start.month, start.day, tzinfo=timezone.utc).timestamp() * 1000)
    ids: set[str] = set()
    params = {"since": since}
    while True:
        response = session.get(
            f"{DST_BASE_URL}/rest/api/3/worklog/updated", params=params, timeout=60
        )
        response.raise_for_status()
        payload = response.json()
        ids.update(str(item["worklogId"]) for item in payload.get("values", []) if item.get("worklogId"))
        if payload.get("lastPage", True):
            break
        params["since"] = payload.get("until", params["since"])

    worklogs_by_id: dict[str, dict[str, Any]] = {}
    ids_list = sorted(ids)
    for index in range(0, len(ids_list), 1000):
        response = session.post(
            f"{DST_BASE_URL}/rest/api/3/worklog/list",
            params={"expand": "properties"},
            json={"ids": ids_list[index:index + 1000]},
            timeout=60,
        )
        response.raise_for_status()
        for worklog in response.json():
            worklog_id = str(worklog.get("id", ""))
            if worklog_id:
                worklogs_by_id[worklog_id] = worklog

    # O feed acima é baseado na data de atualização. A busca por worklogDate
    # também encontra apontamentos futuros que tenham sido criados antes do período.
    issue_keys: set[str] = set()
    next_page_token: str | None = None
    jql = f'worklogDate >= "{start.isoformat()}" AND worklogDate <= "{end.isoformat()}"'
    while True:
        search_body: dict[str, Any] = {
            "jql": jql,
            "fields": ["key"],
            "maxResults": 100,
        }
        if next_page_token:
            search_body["nextPageToken"] = next_page_token
        response = session.post(
            f"{DST_BASE_URL}/rest/api/3/search/jql",
            json=search_body,
            timeout=60,
        )
        response.raise_for_status()
        payload = response.json()
        issue_keys.update(
            str(issue["key"])
            for issue in payload.get("issues", [])
            if issue.get("key")
        )
        next_page_token = payload.get("nextPageToken")
        if not next_page_token:
            break

    query_start = int(
        datetime.combine(
            start - timedelta(days=2), datetime.min.time(), tzinfo=timezone.utc
        ).timestamp() * 1000
    )
    query_end = int(
        datetime.combine(
            end + timedelta(days=2), datetime.min.time(), tzinfo=timezone.utc
        ).timestamp() * 1000
    )
    for issue_key in sorted(issue_keys):
        start_at = 0
        while True:
            response = session.get(
                f"{DST_BASE_URL}/rest/api/3/issue/{issue_key}/worklog",
                params={
                    "expand": "properties",
                    "startAt": start_at,
                    "maxResults": 5000,
                    "startedAfter": query_start,
                    "startedBefore": query_end,
                },
                timeout=60,
            )
            response.raise_for_status()
            payload = response.json()
            page = payload.get("worklogs", [])
            for worklog in page:
                worklog_id = str(worklog.get("id", ""))
                if worklog_id:
                    worklogs_by_id[worklog_id] = worklog
            start_at += len(page)
            if not page or start_at >= int(payload.get("total", len(page))):
                break

    for worklog in worklogs_by_id.values():
        if worklog.get("author", {}).get("accountId") != account_id:
            continue
        day = _record_date(worklog)
        if not day or not (start.isoformat() <= day <= end.isoformat()):
            continue
        daily_seconds[day] += int(worklog.get("timeSpentSeconds") or 0)
        for prop in worklog.get("properties", []):
            if prop.get("key") != SYNC_PROPERTY_KEY:
                continue
            value = prop.get("value")
            marker = value.get("marker") if isinstance(value, dict) else value
            if marker:
                markers.add(str(marker))
    return markers, dict(daily_seconds)


def post_destination_worklog(session: requests.Session, item: dict[str, Any]) -> requests.Response:
    body = {
        "timeSpentSeconds": int(item.get("seconds") or item.get("timeSpentSeconds") or 0),
        "started": item["started"],
        "comment": {
            "type": "doc", "version": 1,
            "content": [{"type": "paragraph", "content": [{"type": "text", "text": item["description"]}]}],
        },
        "properties": [{
            "key": SYNC_PROPERTY_KEY,
            "value": {
                "marker": item["source_marker"],
                "sourceIssue": item.get("key", ""),
                "sourceUpdated": item.get("updated", ""),
            },
        }],
    }
    return session.post(
        f"{DST_BASE_URL}/rest/api/3/issue/{item['target_issue']}/worklog",
        json=body, timeout=60,
    )


def apply_ready_worklogs(
    session: requests.Session,
    plan: Iterable[dict[str, Any]],
    existing_daily_seconds: dict[str, int] | None = None,
    limit_seconds: int = DAILY_LIMIT_SECONDS,
) -> tuple[int, int]:
    """Grava itens elegíveis, atualizando o total somente após cada sucesso."""
    running = defaultdict(int, {
        str(day): int(seconds)
        for day, seconds in (existing_daily_seconds or {}).items()
    })
    failures = 0
    blocked = 0
    successes = 0
    current_day = ""

    print("")
    print("🚀 Iniciando gravação dos apontamentos no Jira interno...")

    for item in plan:
        if item.get("status") != "ready":
            continue
        day = _record_date(item)
        if day != current_day:
            current_day = day
            print("")
            print("=" * 72)
            print(f"📆 {_format_day_label(day)}")
            print("=" * 72)
        projected = running[day] + int(item.get("seconds") or item.get("timeSpentSeconds") or 0)
        print(
            f"   🔹 [{item.get('key', '-')}] → [{item.get('target_issue', '-')}] "
            f"| {seconds_to_hm(int(item.get('seconds') or item.get('timeSpentSeconds') or 0))}"
        )
        if projected > limit_seconds:
            blocked += 1
            print(
                f"      ⚠️  Não registrado: provável duplicação. "
                f"Total atual {seconds_to_hm(running[day])}; "
                f"projetado {seconds_to_hm(projected)} (> 8h).",
                file=sys.stderr,
            )
            continue

        response = post_destination_worklog(session, item)
        if response.ok:
            running[day] = projected
            successes += 1
            print(
                f"      ✅ Registrado com sucesso: [{item['key']}] → [{item['target_issue']}]"
            )
            print(f"      ↳  Total do dia: {seconds_to_hm(running[day])}")
        else:
            failures += 1
            print(
                f"      ❌ Erro ao registrar [{item['key']}] → "
                f"[{item['target_issue']}]: HTTP {response.status_code}",
                file=sys.stderr,
            )

    print("")
    print("=" * 72)
    print(
        f"📊 Processamento concluído: {successes} registrado(s), "
        f"{blocked} bloqueado(s) e {failures} erro(s)."
    )
    print("=" * 72)
    if blocked:
        print(
            f"ATENÇÃO: {blocked} apontamento(s) não foram gravados por ultrapassarem 8h no dia.",
            file=sys.stderr,
        )
    return failures, blocked


def load_records(path: str | Path) -> list[dict[str, Any]]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ValueError("O arquivo JSON deve conter uma lista de worklogs.")
    return [dict(item) for item in payload]


def _format_day_label(day_text: str) -> str:
    weekdays = (
        "Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira",
        "Sexta-feira", "Sábado", "Domingo",
    )
    day = date.fromisoformat(day_text)
    return f"{weekdays[day.weekday()]}, {day.strftime('%d/%m/%Y')}"


def _console_status(item: dict[str, Any], is_application: bool) -> str:
    status = item.get("status")
    if status == "ready":
        return "✅ Será gravado" if is_application else "✅ Pronto para gravar"
    if status == "already_synced":
        return "⏭️  Já sincronizado"
    if status == "probable_duplicate":
        return "⚠️  PROVÁVEL DUPLICAÇÃO (> 8h)"
    if status == "unmapped":
        return "❌ Sem mapeamento — não será gravado"
    return f"⚠️  {status or 'Situação não informada'}"


def print_plan(plan: Iterable[dict[str, Any]], mode: str) -> None:
    rows = list(plan)
    is_application = mode.startswith("APLICAÇÃO")
    report_name = "APLICAÇÃO" if is_application else "SIMULAÇÃO"
    total_seconds = sum(int(item.get("seconds", 0)) for item in rows)
    by_day: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in rows:
        by_day[_record_date(item)].append(item)

    print("")
    print("=" * 72)
    print(f"📋 RELATÓRIO DE {report_name}")
    print("=" * 72)
    print(f"   Modo         : {mode}")
    print(f"   Apontamentos : {len(rows)}")
    print(f"   Horas        : {seconds_to_hm(total_seconds)}")

    if not rows:
        print("")
        print("   Nenhum apontamento encontrado no período.")
        print("=" * 72)
        return

    for day_text in sorted(by_day):
        day_items = by_day[day_text]
        daily_total = max(
            (int(item.get("daily_total_seconds", 0)) for item in day_items),
            default=0,
        )
        print("")
        print("=" * 72)
        print(
            f"📆 {_format_day_label(day_text)} | "
            f"Total do dia: {seconds_to_hm(daily_total)}"
        )
        print("=" * 72)

        for item in day_items:
            target = item.get("target_issue") or "SEM MAPEAMENTO"
            summary = _record_summary(item) or "Sem título"
            print(f"   🔹 [{item.get('key', '-')}] {summary}")
            print(f"      ↳  Origem   : {item.get('comment') or 'Sem comentário'}")
            print(f"      ↳  Destino  : [{target}]")
            print(f"      ↳  Horas    : {seconds_to_hm(int(item.get('seconds', 0)))}")
            print("      ↳  Descrição:")
            for line in str(item.get("description") or "-").splitlines():
                print(f"         {line}")
            print(f"      ↳  Situação : {_console_status(item, is_application)}")
            if item.get("status") == "probable_duplicate":
                print(
                    f"      ↳  Total atual do dia: "
                    f"{seconds_to_hm(int(item.get('daily_total_seconds', 0)))}"
                )
                print(
                    f"      ↳  Total projetado: "
                    f"{seconds_to_hm(int(item.get('daily_projected_seconds', 0)))}"
                )
            elif "daily_total_seconds" in item:
                print(
                    f"      ↳  Total do dia: "
                    f"{seconds_to_hm(int(item['daily_total_seconds']))}"
                )
            print("")

    ready = sum(item.get("status") == "ready" for item in rows)
    synced = sum(item.get("status") == "already_synced" for item in rows)
    blocked = sum(item.get("status") == "probable_duplicate" for item in rows)
    unmapped = sum(item.get("status") == "unmapped" for item in rows)
    print("=" * 72)
    print(
        f"📊 Resumo: {ready} pronto(s), {synced} já sincronizado(s), "
        f"{blocked} bloqueado(s) e {unmapped} sem mapeamento."
    )
    print("=" * 72)


def _excel_safe_value(value: Any) -> Any:
    """Mantém texto vindo do Jira como texto literal dentro do Excel."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def export_plan_xlsx(
    plan: Iterable[dict[str, Any]],
    output_path: str | Path,
    start: date,
    end: date,
    use_ai: bool,
    ai_model: str = OPENAI_MODEL,
) -> Path:
    """Gera uma planilha legível com o resultado completo da simulação."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:  # pragma: no cover - depende do ambiente do usuário
        raise RuntimeError(
            "A geração da planilha requer openpyxl. Execute: "
            "python3 -m pip install -r requirements.txt"
        ) from exc

    rows = list(plan)
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Simulação"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A10"

    navy = "17365D"
    blue = "1F4E78"
    pale_blue = "D9EAF7"
    green = "E2F0D9"
    amber = "FFF2CC"
    red = "FCE4D6"
    gray = "E7E6E6"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="D9E1F2")

    worksheet["A2"] = "Simulação de apontamentos do Jira"
    worksheet["A2"].font = Font(name="Arial", size=16, bold=True, color=navy)
    worksheet["A2"].alignment = Alignment(vertical="center")
    for cell in worksheet[2][0:12]:
        cell.border = Border(bottom=Side(style="medium", color=blue))
    worksheet.row_dimensions[2].height = 28

    worksheet["A3"] = f"Período: {start.strftime('%d/%m/%Y')} a {end.strftime('%d/%m/%Y')}"
    worksheet["A3"].font = Font(name="Arial", size=10, color=navy)
    worksheet["A4"] = (
        f"IA: ativada ({ai_model})" if use_ai else "IA: desativada — descrições geradas por regras locais"
    )
    worksheet["A4"].font = Font(name="Arial", size=10, italic=True, color="44546A")

    summary = [
        ("Apontamentos", len(rows), pale_blue),
        ("Horas", sum(int(item.get("seconds", 0)) for item in rows) / 3600, pale_blue),
        ("Prontos", sum(item.get("status") == "ready" for item in rows), green),
        ("Já sincronizados", sum(item.get("status") == "already_synced" for item in rows), gray),
        ("Prováveis duplicações", sum(item.get("status") == "probable_duplicate" for item in rows), amber),
        ("Sem mapeamento", sum(item.get("status") == "unmapped" for item in rows), red),
    ]
    for index, (label, value, color) in enumerate(summary):
        column = 1 + index * 2
        label_cell = worksheet.cell(6, column, label)
        value_cell = worksheet.cell(7, column, value)
        for row_number in (6, 7):
            left_cell = worksheet.cell(row_number, column)
            right_cell = worksheet.cell(row_number, column + 1)
            left_cell.border = Border(left=thin_gray, top=thin_gray, bottom=thin_gray)
            right_cell.border = Border(right=thin_gray, top=thin_gray, bottom=thin_gray)
            for cell in (left_cell, right_cell):
                cell.fill = PatternFill("solid", fgColor=color)
                cell.alignment = Alignment(horizontal="centerContinuous", vertical="center")
        for cell in (label_cell, value_cell):
            cell.fill = PatternFill("solid", fgColor=color)
        label_cell.font = Font(name="Arial", size=9, bold=True, color=navy)
        value_cell.font = Font(name="Arial", size=14, bold=True, color=navy)
        if label == "Horas":
            value_cell.number_format = '0.00"h"'
    worksheet.row_dimensions[6].height = 21
    worksheet.row_dimensions[7].height = 27

    headers = [
        "Data", "Horas", "Issue origem", "Comentário no cliente", "Título no cliente",
        "Tarefa pai", "Destino", "Status da tarefa", "Descrição que seria gravada",
        "Situação", "Total do dia", "Observação",
    ]
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(9, column, header)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[9].height = 32

    status_labels = {
        "ready": "Pronto para gravar",
        "already_synced": "Já sincronizado",
        "probable_duplicate": "Provável duplicação",
        "unmapped": "Sem mapeamento",
    }
    status_fills = {
        "ready": green,
        "already_synced": gray,
        "probable_duplicate": amber,
        "unmapped": red,
    }
    observations = {
        "ready": "Seria gravado no Jira interno.",
        "already_synced": "Já foi sincronizado por este script; não seria duplicado.",
        "probable_duplicate": "Não seria gravado: o total projetado ultrapassaria 8h.",
        "unmapped": "Sem tarefa interna mapeada; não seria gravado.",
    }

    for row_number, item in enumerate(rows, 10):
        status = str(item.get("status", ""))
        day_text = _record_date(item)
        daily_seconds = (
            item.get("daily_projected_seconds", 0)
            if status == "probable_duplicate"
            else item.get("daily_total_seconds", 0)
        )
        values = [
            date.fromisoformat(day_text) if day_text else None,
            int(item.get("seconds", 0)) / 3600,
            item.get("key", ""),
            item.get("comment", ""),
            _record_summary(item),
            _record_parent(item),
            item.get("target_issue") or "SEM MAPEAMENTO",
            item.get("status_text") or item.get("sourceStatus", ""),
            item.get("description", ""),
            status_labels.get(status, status or "Não informado"),
            int(daily_seconds or 0) / 3600,
            observations.get(status, "Revisar antes de aplicar."),
        ]
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(row_number, column, _excel_safe_value(value))
            cell.font = Font(name="Arial", size=10, color="222222")
            cell.alignment = Alignment(vertical="top", wrap_text=column in {4, 5, 6, 8, 9, 10, 12})
            cell.border = Border(bottom=thin_gray)
        worksheet.cell(row_number, 1).number_format = "dd/mm/yyyy"
        worksheet.cell(row_number, 2).number_format = '0.00"h"'
        worksheet.cell(row_number, 11).number_format = '0.00"h"'
        worksheet.cell(row_number, 10).fill = PatternFill(
            "solid", fgColor=status_fills.get(status, pale_blue)
        )
        worksheet.cell(row_number, 10).font = Font(name="Arial", size=10, bold=True, color=navy)
        wrapped_fields = zip(
            (values[index] for index in (3, 4, 5, 7, 8, 9, 11)),
            (22, 34, 30, 20, 58, 23, 42),
        )
        visual_lines = 1
        for value, approximate_width in wrapped_fields:
            text = str(value or "")
            field_lines = sum(
                max(1, (len(line) + approximate_width - 1) // approximate_width)
                for line in text.splitlines() or [""]
            )
            visual_lines = max(visual_lines, field_lines)
        worksheet.row_dimensions[row_number].height = max(30, min(180, 15 * visual_lines))

    if rows:
        table = Table(displayName="SimulacaoApontamentos", ref=f"A9:L{9 + len(rows)}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        worksheet.add_table(table)
    else:
        worksheet["A10"] = "Nenhum apontamento encontrado no período informado."
        worksheet["A10"].font = Font(name="Arial", italic=True, color="666666")
        worksheet.row_dimensions[10].height = 32

    widths = [12, 10, 15, 22, 34, 30, 15, 20, 58, 23, 14, 42]
    for column, width in enumerate(widths, 1):
        worksheet.column_dimensions[chr(64 + column)].width = width
    worksheet.auto_filter.ref = f"A9:L{max(9 + len(rows), 10)}"
    worksheet.print_title_rows = "1:9"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_area = f"A2:L{max(9 + len(rows), 10)}"

    workbook.properties.title = "Simulação de apontamentos do Jira"
    workbook.properties.subject = "Prévia dos apontamentos que seriam gravados no Jira interno"
    workbook.properties.creator = "Jira Sync"
    workbook.save(output)
    return output


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Sincroniza worklogs com descrições no formato solicitado pelo RH.")
    parser.add_argument("--start", help="Data inicial YYYY-MM-DD")
    parser.add_argument("--end", help="Data final YYYY-MM-DD")
    parser.add_argument("--weeks-back", type=int, default=WEEKS_BACK)
    parser.add_argument("--input-json", help="Exportação JSON para simulação/reprocessamento")
    parser.add_argument("--apply", action="store_true", help="Grava no Jira destino; sem esta opção é simulação")
    parser.add_argument(
        "--output-xlsx",
        help="Salva a simulação em uma planilha XLSX, sem exibir todos os itens no terminal",
    )
    ai_group = parser.add_mutually_exclusive_group()
    ai_group.add_argument("--ai", dest="use_ai", action="store_true", help="Melhora as descrições com IA")
    ai_group.add_argument("--no-ai", dest="use_ai", action="store_false", help="Usa apenas as regras locais")
    parser.set_defaults(use_ai=None)
    parser.add_argument(
        "--ai-model", default=OPENAI_MODEL,
        help=f"Modelo da OpenAI (padrão: {OPENAI_MODEL})",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.apply and args.output_xlsx:
        raise SystemExit("--output-xlsx não pode ser usado com --apply.")
    use_ai = resolve_ai_choice(args.use_ai)
    if bool(args.start) != bool(args.end):
        raise SystemExit("Informe --start e --end juntos.")
    if args.start:
        start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)
    else:
        start, end = get_period_range(args.weeks_back)
    if end < start:
        raise SystemExit("A data final não pode ser anterior à inicial.")

    records = load_records(args.input_json) if args.input_json else fetch_source_records(start, end)
    records = [item for item in records if start.isoformat() <= _record_date(item) <= end.isoformat()]

    destination_session: requests.Session | None = None
    existing: set[str] = set()
    existing_daily_seconds: dict[str, int] = {}
    destination_state_available = False
    if DST_EMAIL and DST_API_TOKEN:
        destination_session = _session(DST_EMAIL, DST_API_TOKEN)
        try:
            destination_account = DST_ACCOUNT_ID or _account_id(destination_session, DST_BASE_URL)
            existing, existing_daily_seconds = get_destination_state(
                destination_session, destination_account, start, end
            )
            destination_state_available = True
        except requests.RequestException:
            if args.apply:
                raise
            print(
                "AVISO: não foi possível consultar o Jira interno; os totais exibidos "
                "consideram somente os itens desta simulação.",
                file=sys.stderr,
            )
    elif args.apply:
        raise RuntimeError("Defina DST_EMAIL e DST_API_TOKEN antes de usar --apply.")
    else:
        print(
            "AVISO: credenciais do Jira interno não configuradas; os totais exibidos "
            "consideram somente os itens desta simulação.",
            file=sys.stderr,
        )

    plan = build_sync_plan(records, existing, SRC_BASE_URL)
    if use_ai:
        plan = enhance_plan_with_ai(plan, OPENAI_API_KEY, args.ai_model)
        statuses = {item.get("ai_status") for item in plan}
        if "fallback_missing_key" in statuses:
            print(
                "AVISO: OPENAI_API_KEY não configurada; usando as descrições locais.",
                file=sys.stderr,
            )
        elif "fallback_error" in statuses:
            print(
                "AVISO: a IA ficou indisponível; os itens afetados mantiveram as descrições locais.",
                file=sys.stderr,
            )
        elif "fallback_invalid" in statuses:
            print(
                "AVISO: algumas respostas da IA foram rejeitadas pelas validações e mantiveram a versão local.",
                file=sys.stderr,
            )
    mode = "APLICAÇÃO" if args.apply else "SIMULAÇÃO — nenhuma gravação será feita"
    mode += " | IA ativada" if use_ai else " | IA desativada"
    mode += (
        " | horas existentes do destino consideradas"
        if destination_state_available
        else " | totais apenas desta execução"
    )
    preview_plan = apply_daily_limit(plan, existing_daily_seconds)
    if args.output_xlsx:
        output = export_plan_xlsx(
            preview_plan, args.output_xlsx, start, end, use_ai, args.ai_model
        )
        print(f"Planilha gerada: {output}")
    else:
        print_plan(preview_plan, mode)
    if not args.apply:
        return 0

    assert destination_session is not None
    failures, _blocked = apply_ready_worklogs(
        destination_session, plan, existing_daily_seconds
    )
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
